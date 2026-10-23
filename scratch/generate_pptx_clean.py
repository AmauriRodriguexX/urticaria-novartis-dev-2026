import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

COLOR_INK = RGBColor(34, 56, 63)         # #22383f
COLOR_ACCENT = RGBColor(224, 88, 77)     # #e0584d
COLOR_TEAL = RGBColor(0, 129, 168)       # #0081a8
COLOR_BG_CARD = RGBColor(247, 246, 242)  # #f7f6f2
COLOR_LINE = RGBColor(220, 216, 201)     # #dcd8c9
COLOR_MUTED = RGBColor(90, 106, 109)     # #5a6a6d
COLOR_WHITE = RGBColor(255, 255, 255)

SECTIONS_DATA = [
    {
        "code": "SECCIÓN 01",
        "name": "HOME Y VARIANTES CONTEXTUALES",
        "desc": "Portadas de inicio dinámicas adaptadas según clima y hora (Día, Calor, Frío, Madrugada y Noche).",
        "screens": [
            {
                "file": "01_home_dia.png",
                "section": "SECCIÓN: HOME",
                "title": "Home Principal — Escenario Día (General)"
            },
            {
                "file": "02_home_calor.png",
                "section": "SECCIÓN: HOME",
                "title": "Home Principal — Escenario Calor (Temperatura Alta / Sol)"
            },
            {
                "file": "03_home_frio.png",
                "section": "SECCIÓN: HOME",
                "title": "Home Principal — Escenario Frío (Invierno / Aire Helado)"
            },
            {
                "file": "04_home_madrugada.png",
                "section": "SECCIÓN: HOME",
                "title": "Home Principal — Escenario Madrugada (Modo Oscuro 2am - 5am)"
            },
            {
                "file": "05_home_noche_calor.png",
                "section": "SECCIÓN: HOME",
                "title": "Home Principal — Escenario Noche Cálida"
            }
        ]
    },
    {
        "code": "SECCIÓN 02",
        "name": "LANDINGS INFORMATIVAS Y CLÍNICAS",
        "desc": "Páginas completas de contenidos médicos, tratamientos, información institucional y error 404.",
        "screens": [
            {
                "file": "06_sintomas.png",
                "section": "SECCIÓN: LANDINGS INFORMATIVAS",
                "title": "Síntomas de la Urticaria Crónica Espontánea (/sintomas/)"
            },
            {
                "file": "07_tratamiento.png",
                "section": "SECCIÓN: LANDINGS INFORMATIVAS",
                "title": "Tratamientos para la UCE (/tratamiento/)"
            },
            {
                "file": "08_sobre_nosotros.png",
                "section": "SECCIÓN: LANDINGS INFORMATIVAS",
                "title": "Sobre Nosotros — Más que Urticaria (/sobre-nosotros/)"
            },
            {
                "file": "09_error_404.png",
                "section": "SECCIÓN: LANDINGS INFORMATIVAS",
                "title": "Página de Error 404 (Ruta no encontrada)"
            }
        ]
    },
    {
        "code": "SECCIÓN 03",
        "name": "RECURSOS Y HERRAMIENTAS",
        "desc": "Hub de selección de herramientas, bitácora del diario de síntomas, comunidad y biblioteca visual.",
        "screens": [
            {
                "file": "10_recursos_hub.png",
                "section": "SECCIÓN: RECURSOS Y HERRAMIENTAS",
                "title": "Hub de Recursos y Herramientas (/recursos-herramientas/)"
            },
            {
                "file": "11_diario_sintomas.png",
                "section": "SECCIÓN: RECURSOS Y HERRAMIENTAS",
                "title": "Diario de Síntomas — Bitácora Interactiva (/diario/)"
            },
            {
                "file": "12_testimonios_comunidad.png",
                "section": "SECCIÓN: RECURSOS Y HERRAMIENTAS",
                "title": "Testimonios y Comunidad — 'No lo estás imaginando' (/testimonios/)"
            },
            {
                "file": "19_biblioteca_modal.png",
                "section": "SECCIÓN: RECURSOS Y HERRAMIENTAS",
                "title": "Biblioteca Visual — Ficha de Fototipo por Tono de Piel (Modal Lightbox)"
            }
        ]
    },
    {
        "code": "SECCIÓN 04",
        "name": "QUIZ ORIENTATIVO & URGENCIAS",
        "desc": "Flujo interactivo completo del cuestionario clínico (Q0 a Q19), alerta 911 y pantallas de resultados.",
        "screens": [
            {
                "file": "18_quiz_intro_landing.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Landing Introductoria del Test de Síntomas (/test/)"
            },
            {
                "file": "24_quiz_q0_nombre.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 0 — Nombre / Personalización (Opcional)"
            },
            {
                "file": "25_quiz_q1_sintomas.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 1 — Síntomas Experimentados (Ronchas, Picazón, Hinchazón)"
            },
            {
                "file": "20_quiz_q1b_descarte_urgencia.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 1b — Descarte de Urgencia Crítica (Labios, lengua, garganta)"
            },
            {
                "file": "21_quiz_urgencia_911.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "🚨 Pantalla de Urgencia Médica (Alerta Inmediata 911)"
            },
            {
                "file": "22_quiz_modal_directorio.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Modal Embebido de Especialistas (Buscador dentro del Quiz)"
            },
            {
                "file": "26_quiz_q2_duracion.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 2 — Tiempo con Síntomas (< 1 sem hasta > 6 meses)"
            },
            {
                "file": "27_quiz_q3_ronchas.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 3 — Cantidad de Ronchas en 24h (Escala UAS7)"
            },
            {
                "file": "28_quiz_q4_picazon.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 4 — Intensidad de la Comezón (Escala UAS7)"
            },
            {
                "file": "29_quiz_q5_hinchazon.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 5 — Severidad de la Hinchazón / Angioedema"
            },
            {
                "file": "30_quiz_q6_calidad_vida.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 6 — Áreas de Calidad de Vida Afectadas (CU-Q2oL)"
            },
            {
                "file": "31_quiz_q7_actividades.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 7 — Interferencia en Actividades Diarias"
            },
            {
                "file": "32_quiz_q8_sociales.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 8 — Afectación en Relaciones Familiares y Sociales"
            },
            {
                "file": "33_quiz_q8b_sexual.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 8b — Afectación en Vida Sexual e Intimidad"
            },
            {
                "file": "34_quiz_q9_vestimenta.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 9 — Influencia en la Ropa y Vestimenta que Usas"
            },
            {
                "file": "35_quiz_q10_trabajo.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 10 — Afectación en Trabajo o Estudios"
            },
            {
                "file": "36_quiz_q11_sueno.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 11 — Afectación en Calidad de Sueño (Insomnio nocturno)"
            },
            {
                "file": "37_quiz_q12_detonantes.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 12 — Posibles Desencadenantes (Calor, frío, estrés, sudor)"
            },
            {
                "file": "38_quiz_q13_horario.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 13 — Momento del Día de Mayor Intensidad"
            },
            {
                "file": "39_quiz_q14_consulta_previa.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 14 — ¿Has Consultado Previamente a un Médico?"
            },
            {
                "file": "40_quiz_q15_frecuencia_consulta.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 15 — Frecuencia de Consulta Médica"
            },
            {
                "file": "41_quiz_q16_especialista.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 16 — Especialistas Consultados (Dermatólogo, Alergólogo, General)"
            },
            {
                "file": "42_quiz_q17_tratamiento_previo.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 17 — ¿Has Tomado o Aplicado Tratamiento Previo?"
            },
            {
                "file": "43_quiz_q18_tipo_tratamiento.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 18 — Tipo de Tratamiento Recibido (Básico, Intermedio, Avanzado)"
            },
            {
                "file": "44_quiz_q19_efectividad.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Paso 19 — Efectividad Percibida del Tratamiento"
            },
            {
                "file": "45_quiz_resultado_reporte_completo.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Pantalla de Resultados — Reporte Clínico Completo (Score Ben, PDF y CTA)"
            },
            {
                "file": "23_quiz_resultado_asintomatico.png",
                "section": "SECCIÓN: QUIZ ORIENTATIVO",
                "title": "Pantalla de Resultados — Ruta Asintomática (Ningún síntoma en Paso 1)"
            }
        ]
    },
    {
        "code": "SECCIÓN 05",
        "name": "DIRECTORIO MÉDICO & BLOG EDITORIAL",
        "desc": "Landing del directorio de especialistas en México y artículos de lectura educativa.",
        "screens": [
            {
                "file": "13_directorio_landing.png",
                "section": "SECCIÓN: DIRECTORIO & BLOG",
                "title": "Directorio Médico de Especialistas — Landing Principal (/directorio-medico/)"
            },
            {
                "file": "14_blog_feed.png",
                "section": "SECCIÓN: DIRECTORIO & BLOG",
                "title": "Blog Editorial — Catálogo de Temas (/blog/)"
            },
            {
                "file": "15_blog_articulo_1.png",
                "section": "SECCIÓN: DIRECTORIO & BLOG",
                "title": "Artículo Blog 1: ¿Qué significa UCE? (/blog/que-significa-uce/)"
            },
            {
                "file": "16_blog_articulo_2.png",
                "section": "SECCIÓN: DIRECTORIO & BLOG",
                "title": "Artículo Blog 2: Guía para preparar tu consulta (/blog/preparar-consulta/)"
            },
            {
                "file": "17_blog_articulo_3.png",
                "section": "SECCIÓN: DIRECTORIO & BLOG",
                "title": "Artículo Blog 3: Mitos sobre dietas restrictivas (/blog/dietas-restrictivas-uce/)"
            }
        ]
    }
]

def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]
    screenshot_dir = os.path.abspath("screenshots_viewport")

    # -------------------------------------------------------------
    # SLIDE 1: PORTADA PRINCIPAL
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_INK
    bg1.line.fill.background()

    bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.25))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = COLOR_ACCENT
    bar1.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "NOVARTIS MÉXICO · MÁS QUE URTICARIA"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT

    p1 = tf1.add_paragraph()
    p1.text = "Catálogo Visual de Pantallas & Flujos"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_before = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Capturas completas de la primera sección visible (viewport laptop 1366 × 768), organizadas por secciones temáticas:\nHome, Landings Informativas, Recursos, Quiz Orientativo y Directorio Médico."
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(200, 215, 220)
    p2.space_before = Pt(16)

    p3 = tf1.add_paragraph()
    p3.text = "Octubre 2026 | Responsables: Amauri Rodríguez Santos / Edward Alarcón"
    p3.font.size = Pt(12)
    p3.font.color.rgb = RGBColor(160, 180, 185)
    p3.space_before = Pt(28)

    # -------------------------------------------------------------
    # SLIDE 2: ÍNDICE GENERAL POR SECCIONES
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    
    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(1.0))
    tf2 = tb2.text_frame
    p_i0 = tf2.paragraphs[0]
    p_i0.text = "ESTRUCTURA DEL SITIO WEB"
    p_i0.font.size = Pt(12)
    p_i0.font.bold = True
    p_i0.font.color.rgb = COLOR_ACCENT

    p_i1 = tf2.add_paragraph()
    p_i1.text = "Índice de Secciones y Pantallas (45 Pantallas)"
    p_i1.font.size = Pt(24)
    p_i1.font.bold = True
    p_i1.font.color.rgb = COLOR_INK

    card_w = Inches(2.25)
    card_h = Inches(5.0)
    card_gap = Inches(0.18)
    start_x = Inches(0.8)
    start_y = Inches(1.7)

    colors = [COLOR_ACCENT, COLOR_TEAL, COLOR_MUTED, COLOR_ACCENT, COLOR_TEAL]

    for i, sec in enumerate(SECTIONS_DATA):
        cx = start_x + i * (card_w + card_gap)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, start_y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_LINE
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True

        cp0 = ctf.paragraphs[0]
        cp0.text = sec["code"]
        cp0.font.size = Pt(11)
        cp0.font.bold = True
        cp0.font.color.rgb = colors[i]

        cp1 = ctf.add_paragraph()
        cp1.text = sec["name"]
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = COLOR_INK
        cp1.space_before = Pt(6)

        cp2 = ctf.add_paragraph()
        cp2.text = f"{len(sec['screens'])} Vistas"
        cp2.font.size = Pt(12)
        cp2.font.bold = True
        cp2.font.color.rgb = colors[i]
        cp2.space_before = Pt(8)

        cp3 = ctf.add_paragraph()
        cp3.text = sec["desc"]
        cp3.font.size = Pt(9.5)
        cp3.font.color.rgb = COLOR_MUTED
        cp3.space_before = Pt(10)

    # -------------------------------------------------------------
    # SECCIONES: SEPARADOR + PANTALLAS A TAMAÑO COMPLETO
    # -------------------------------------------------------------
    global_counter = 1

    for s_idx, sec in enumerate(SECTIONS_DATA, start=1):
        # 1. Slide Separador de Sección
        slide_sep = prs.slides.add_slide(blank_layout)
        bg_sep = slide_sep.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
        bg_sep.fill.solid()
        bg_sep.fill.fore_color.rgb = COLOR_INK
        bg_sep.line.fill.background()

        bar_sep = slide_sep.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.4), SLIDE_HEIGHT)
        bar_sep.fill.solid()
        bar_sep.fill.fore_color.rgb = COLOR_ACCENT
        bar_sep.line.fill.background()

        tb_sep = slide_sep.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.5), Inches(3.5))
        tf_sep = tb_sep.text_frame
        tf_sep.word_wrap = True

        sep0 = tf_sep.paragraphs[0]
        sep0.text = f"{sec['code']} · DIVISIÓN DE VISTAS"
        sep0.font.size = Pt(16)
        sep0.font.bold = True
        sep0.font.color.rgb = COLOR_ACCENT

        sep1 = tf_sep.add_paragraph()
        sep1.text = sec["name"]
        sep1.font.size = Pt(36)
        sep1.font.bold = True
        sep1.font.color.rgb = COLOR_WHITE
        sep1.space_before = Pt(10)

        sep2 = tf_sep.add_paragraph()
        sep2.text = f"{sec['desc']}\nTotal de pantallas en este bloque: {len(sec['screens'])} vistas a pantalla completa."
        sep2.font.size = Pt(16)
        sep2.font.color.rgb = RGBColor(200, 215, 220)
        sep2.space_before = Pt(14)

        # 2. Slides de Pantallas de la Sección
        for sc in sec["screens"]:
            img_path = os.path.join(screenshot_dir, sc["file"])
            if not os.path.exists(img_path):
                print(f"Aviso: archivo no encontrado {img_path}")
                continue

            slide = prs.slides.add_slide(blank_layout)

            # ENCABEZADO MÍNIMO DE 2 RENGLONES (como pidió el usuario)
            hdr_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.2), Inches(12.13), Inches(0.75))
            htf = hdr_box.text_frame
            htf.word_wrap = True

            # RENGLÓN 1: SECCIÓN
            h_line1 = htf.paragraphs[0]
            h_line1.text = f"{sc['section']}  ·  VISTA {global_counter:02d} DE 45"
            h_line1.font.size = Pt(11)
            h_line1.font.bold = True
            h_line1.font.color.rgb = COLOR_ACCENT

            # RENGLÓN 2: VISTA
            h_line2 = htf.add_paragraph()
            h_line2.text = sc["title"]
            h_line2.font.size = Pt(17)
            h_line2.font.bold = True
            h_line2.font.color.rgb = COLOR_INK

            # CAPTURA A PANTALLA COMPLETA (ocupa casi toda la diapositiva, sin comentarios laterales)
            max_w = Inches(12.13)
            max_h = Inches(6.25)
            origin_x = Inches(0.6)
            origin_y = Inches(1.05)

            with Image.open(img_path) as im:
                orig_w, orig_h = im.size

            img_ratio = orig_w / orig_h
            box_ratio = max_w / max_h

            if img_ratio > box_ratio:
                final_w = max_w
                final_h = max_w / img_ratio
            else:
                final_h = max_h
                final_w = max_h * img_ratio

            pos_x = origin_x + (max_w - final_w) / 2
            pos_y = origin_y + (max_h - final_h) / 2

            # Marco sutil
            bg_frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pos_x - Pt(2), pos_y - Pt(2), final_w + Pt(4), final_h + Pt(4))
            bg_frame.fill.solid()
            bg_frame.fill.fore_color.rgb = COLOR_LINE
            bg_frame.line.fill.background()

            slide.shapes.add_picture(img_path, pos_x, pos_y, width=final_w, height=final_h)

            global_counter += 1

    output_file = "Novartis_Mas_Que_Urticaria_Catalogo_Pantallas.pptx"
    prs.save(output_file)
    print(f"Presentacion generada con exito: {os.path.abspath(output_file)}")

if __name__ == "__main__":
    build_presentation()
