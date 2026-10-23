import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette (Novartis / Más que Urticaria Design System)
COLOR_INK = RGBColor(34, 56, 63)         # #22383f
COLOR_ACCENT = RGBColor(224, 88, 77)     # #e0584d (Coral)
COLOR_TEAL = RGBColor(0, 129, 168)       # #0081a8
COLOR_BG_CARD = RGBColor(247, 246, 242)  # #f7f6f2
COLOR_LINE = RGBColor(220, 216, 201)     # #dcd8c9
COLOR_MUTED = RGBColor(90, 106, 109)     # #5a6a6d
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_AMBER = RGBColor(217, 119, 6)

SCREENS_ORDERED = [
    # =========================================================
    # SECCIÓN 1: HOME Y VARIANTES DE CONTEXTO
    # =========================================================
    {
        "file": "01_home_dia.png",
        "section": "SECCIÓN 1: HOME & CONTEXTO",
        "title": "Home Principal — Escenario Día (General)",
        "route": "URL: /?vista=dia",
        "type": "Landing Web Full-Page",
        "desc": "Vista predeterminada durante el día. Hero con mensaje principal 'Haz de la urticaria la parte más pequeña de tu vida', CTA principal al Quiz orientativo, buscador de especialistas, infografía comparativa de piel, ruta de 3 pasos y testimonios destacados.",
        "evaluar": "Jerarquía visual del Hero, claridad del copy inicial, visibilidad del CTA hacia el test, legibilidad del esquema informativo de habones."
    },
    {
        "file": "02_home_calor.png",
        "section": "SECCIÓN 1: HOME & CONTEXTO",
        "title": "Home — Escenario Calor (Temperatura Alta / Sol)",
        "route": "URL: /?vista=calor",
        "type": "Landing Web Full-Page (Variante)",
        "desc": "Activada automáticamente cuando la API de clima detecta temperaturas elevadas en la ciudad del usuario. Paleta cálida, mensaje enfocado en cómo el calor y sudor despiertan la piel.",
        "evaluar": "Tono empático ante brotes por calor, contraste visual con la paleta cálida, coherencia con guías de UCE inducible por calor/sudor."
    },
    {
        "file": "03_home_frio.png",
        "section": "SECCIÓN 1: HOME & CONTEXTO",
        "title": "Home — Escenario Frío (Invierno / Aire Helado)",
        "route": "URL: /?vista=frio",
        "type": "Landing Web Full-Page (Variante)",
        "desc": "Activada en temporadas frías. Paleta gélida/azulada, mensaje sobre irritación por resequedad ambiental y aclaración sobre urticaria por frío.",
        "evaluar": "Diferenciación adecuada entre UCE y urticaria por frío (a frigore), tono no alarmista."
    },
    {
        "file": "04_home_madrugada.png",
        "section": "SECCIÓN 1: HOME & CONTEXTO",
        "title": "Home — Escenario Madrugada (Insomnio 2am - 5am)",
        "route": "URL: /?vista=madrugada",
        "type": "Landing Web Full-Page (Variante)",
        "desc": "Modo oscuro automático cuando el usuario navega de madrugada. Mensaje de acompañamiento '¿No puedes dormir? La comezón aparece en silencio. No estás solo/a'.",
        "evaluar": "Impacto psicológico y contención emocional, accesibilidad en modo nocturno, descanso visual."
    },
    {
        "file": "05_home_noche_calor.png",
        "section": "SECCIÓN 1: HOME & CONTEXTO",
        "title": "Home — Escenario Noche Cálida",
        "route": "URL: /?vista=noche-calor",
        "type": "Landing Web Full-Page (Variante)",
        "desc": "Variante nocturna en zonas con clima caluroso. Combina el modo oscuro con acentos cálidos y copy enfocado en el bochorno nocturno.",
        "evaluar": "Legibilidad del texto sobre fondo oscuro, llamada a la calma antes de dormir."
    },

    # =========================================================
    # SECCIÓN 2: LANDINGS INFORMATIVAS Y CLÍNICAS
    # =========================================================
    {
        "file": "06_sintomas.png",
        "section": "SECCIÓN 2: LANDINGS INFORMATIVAS",
        "title": "Síntomas de la Urticaria Crónica Espontánea",
        "route": "URL: /sintomas/",
        "type": "Landing Informativa Full-Page",
        "desc": "Explicación de los signos clínicos (habones, prurito, angioedema), tabla comparativa accesible entre UCE, urticaria aguda y urticaria inducible, perfiles por edad/sexo y FAQ.",
        "evaluar": "Revisión médica de las diferencias en la tabla comparativa, precisión de términos dermatológicos, mensaje de urgencia si hay angioedema laríngeo."
    },
    {
        "file": "07_tratamiento.png",
        "section": "SECCIÓN 2: LANDINGS INFORMATIVAS",
        "title": "Tratamientos para la Urticaria Crónica",
        "route": "URL: /tratamiento/",
        "type": "Landing Informativa Full-Page",
        "desc": "Orientación general sobre abordaje médico, función de mastocitos e histamina, opciones de medicamentos (antihistamínicos, corticoides, terapias biológicas dirigidas), qué preguntar al especialista y medidas cotidianas.",
        "evaluar": "Cumplimiento normativo (no prescripción directa, lenguaje neutral), aval de opciones terapéuticas, diferenciación entre dermatología y alergología."
    },
    {
        "file": "08_sobre_nosotros.png",
        "section": "SECCIÓN 2: LANDINGS INFORMATIVAS",
        "title": "Sobre Nosotros — Más que Urticaria",
        "route": "URL: /sobre-nosotros/",
        "type": "Landing Institucional Full-Page",
        "desc": "Página institucional que transparenta la naturaleza informativa del sitio, principios editoriales de claridad y lenguaje accesible, y pie con patrocinio de Novartis.",
        "evaluar": "Declaraciones institucionales y legales autorizadas por Novartis, aviso de exención de responsabilidad médica."
    },
    {
        "file": "09_error_404.png",
        "section": "SECCIÓN 2: LANDINGS INFORMATIVAS",
        "title": "Página de Error 404 (Ruta no disponible)",
        "route": "URL: /404.html",
        "type": "Página de Sistema",
        "desc": "Estado de respaldo ante enlaces rotos o secciones aún no publicadas, guiando al visitante de regreso al Inicio de forma amigable.",
        "evaluar": "Claridad de navegación de retorno, tono respetuoso."
    },

    # =========================================================
    # SECCIÓN 3: RECURSOS Y HERRAMIENTAS
    # =========================================================
    {
        "file": "10_recursos_hub.png",
        "section": "SECCIÓN 3: RECURSOS & HERRAMIENTAS",
        "title": "Hub de Recursos y Herramientas",
        "route": "URL: /recursos-herramientas/",
        "type": "Hub de Navegación",
        "desc": "Guía por intención del paciente: 'Quiero entender mis síntomas' (Test), 'Quiero registrarlos' (Diario) y 'Quiero ver cómo luce en mi piel' (Biblioteca visual).",
        "evaluar": "Comprensión de los 3 caminos, tarjetas destacadas con sus respectivos accesos directos."
    },
    {
        "file": "11_diario_sintomas.png",
        "section": "SECCIÓN 3: RECURSOS & HERRAMIENTAS",
        "title": "Diario de Síntomas — Bitácora Interactiva",
        "route": "URL: /diario/",
        "type": "Herramienta Interactiva",
        "desc": "Bitácora de registro día a día con escala 0-3 (habones y comezón). Incluye función de impresión/guardado en PDF local sin almacenar datos sensibles en servidor.",
        "evaluar": "Privacidad de datos (se recalca que no se envían a base de datos), utilidad de la hoja impresa para el médico en consulta."
    },
    {
        "file": "12_testimonios_comunidad.png",
        "section": "SECCIÓN 3: RECURSOS & HERRAMIENTAS",
        "title": "Testimonios — 'No lo estás imaginando'",
        "route": "URL: /testimonios/",
        "type": "Comunidad & Pacientes",
        "desc": "Espacio de relatos y experiencias compartidas. Muestra casos de muestra etiquetados transparentemente como pendientes de autorización y consentimiento formal.",
        "evaluar": "Cumplimiento ético de testimonios, formato de historias, moderación y aviso de confidencialidad."
    },
    {
        "file": "19_biblioteca_modal.png",
        "section": "SECCIÓN 3: RECURSOS & HERRAMIENTAS",
        "title": "Biblioteca Visual — Ficha de Tono de Piel (Modal)",
        "route": "Modal /recursos-herramientas/#biblioteca",
        "type": "Visor Lightbox Interactivo",
        "desc": "Vista ampliada de fotografías clínicas/educativas clasificadas por fototipos Fitzpatrick (I al VI), ayudando a pacientes con piel morena o clara a reconocer habones.",
        "evaluar": "Sensibilidad y diversidad de tonos de piel, leyenda informativa aclaratoria de que son imágenes de referencia sujetas a revisión médica."
    },

    # =========================================================
    # SECCIÓN 4: QUIZ ORIENTATIVO & URGENCIAS (FLUJO COMPLETO)
    # =========================================================
    {
        "file": "18_quiz_intro_landing.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Landing Introductoria del Test de Síntomas",
        "route": "URL: /test/",
        "type": "Landing Pre-Test",
        "desc": "Explicación previa al inicio del cuestionario: duración estimada (<5 min), aviso explícito de que no diagnostica ni reemplaza consulta médica.",
        "evaluar": "Aviso legal de no-diagnóstico, prominencia del botón de inicio."
    },
    {
        "file": "24_quiz_q0_nombre.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 0 — Personalización / Nombre",
        "route": "Modal Quiz · Paso 0",
        "type": "Input Opcional",
        "desc": "Campo de texto opcional para personalizar el encabezado del reporte para la consulta médica. Puede omitirse y avanzar directamente.",
        "evaluar": "Claridad de que es opcional para proteger la privacidad."
    },
    {
        "file": "25_quiz_q1_sintomas.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 1 — Síntomas Experimentados",
        "route": "Modal Quiz · Paso 1 (Multi-select)",
        "type": "Selección Múltiple",
        "desc": "Filtro inicial de síntomas: Ronchas en la piel, Comezón intensa, Hinchazón (angioedema) o Ninguno de estos.",
        "evaluar": "Descripciones sencillas de cada síntoma, comportamiento exclusivo de la opción 'Ninguno'."
    },
    {
        "file": "20_quiz_q1b_descarte_urgencia.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 1b — Descarte de Urgencia Crítica",
        "route": "Modal Quiz · Paso 1b (Condicional)",
        "type": "Bifurcación de Seguridad",
        "desc": "Aparece únicamente si el usuario seleccionó 'Hinchazón' en el Paso 1. Pregunta si la inflamación es en labios, lengua, garganta o si dificulta respirar.",
        "evaluar": "Redacción de máxima prioridad clínica para evitar retrasos ante anafilaxia."
    },
    {
        "file": "21_quiz_urgencia_911.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "🚨 Pantalla de Urgencia Médica (Alerta 911)",
        "route": "Modal Quiz · Desvío Inmediato",
        "type": "Pantalla de Emergencia",
        "desc": "Interrumpe de inmediato el cuestionario. Muestra alerta de alto impacto con botón directo para marcar al 911 (emergencias médicas) y opción secundaria para ver especialistas.",
        "evaluar": "Efectividad del botón 'Llamar a emergencias (911)', claridad del mensaje de urgencia inmediata."
    },
    {
        "file": "22_quiz_modal_directorio.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Modal Embebido de Directorio Médico",
        "route": "Modal Quiz · phase: 'map'",
        "type": "Herramienta Embebida",
        "desc": "Permite consultar especialistas certificados directamente desde el flujo del test o desde la pantalla de urgencia sin abandonar el modal.",
        "evaluar": "Filtros por especialidad, geolocalización y datos de contacto de especialistas."
    },
    {
        "file": "26_quiz_q2_duracion.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 2 — Tiempo con Síntomas",
        "route": "Modal Quiz · Paso 2",
        "type": "Selección Única",
        "desc": "Evalúa la temporalidad (<1 sem, 1 sem a 1 mes, 1 a 6 meses, >6 meses). Fundamental para distinguir entre brotes agudos y sospecha de UCE (+6 semanas).",
        "evaluar": "Ponderación del umbral de las 6 semanas según criterios clínicos internacionales."
    },
    {
        "file": "27_quiz_q3_ronchas.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 3 — Cantidad de Ronchas (24h)",
        "route": "Modal Quiz · Paso 3 (Condicional)",
        "type": "Severidad Síntomas",
        "desc": "Se activa solo si marcó 'Ronchas'. Escala basada en el UAS7: Ninguna, Leve (<20), Moderada (20-50), Alta (>50 ronchas o áreas confluentes).",
        "evaluar": "Correspondencia con la escala validada UAS7 (Urticaria Activity Score)."
    },
    {
        "file": "28_quiz_q4_picazon.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 4 — Intensidad de la Picazón",
        "route": "Modal Quiz · Paso 4 (Condicional)",
        "type": "Severidad Síntomas",
        "desc": "Se activa solo si marcó 'Comezón'. Escala UAS7 de prurito: Ninguna, Leve, Moderada o Severa (interrumpe actividades o sueño).",
        "evaluar": "Claridad de los niveles de interferencia de la picazón."
    },
    {
        "file": "29_quiz_q5_hinchazon.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 5 — Severidad de la Hinchazón",
        "route": "Modal Quiz · Paso 5 (Condicional)",
        "type": "Severidad Síntomas",
        "desc": "Se activa solo si marcó 'Hinchazón'. Escala AAS (Angioedema Activity Score): Ninguna, Leve, Moderada o Intensa (dolorosa o limita movilidad).",
        "evaluar": "Validación clínica del impacto del angioedema periférico."
    },
    {
        "file": "30_quiz_q6_calidad_vida.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 6 — Áreas de Vida Afectadas",
        "route": "Modal Quiz · Paso 6 (Multi-select)",
        "type": "Eje Calidad de Vida (CU-Q2oL)",
        "desc": "Selección de dimensiones afectadas: Actividades diarias, Vida social, Vida sexual, Elección de vestimenta, Trabajo o estudios, y Calidad de sueño.",
        "evaluar": "Alineación con el cuestionario oficial CU-Q2oL de calidad de vida en urticaria."
    },
    {
        "file": "31_quiz_q7_actividades.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 7 — Interferencia en Actividades Diarias",
        "route": "Modal Quiz · Paso 7 (Subpregunta)",
        "type": "Escala 0 a 3",
        "desc": "Se activa si seleccionó 'Actividades diarias' en Q6. Evalúa impacto en tareas del hogar, compras, etc. (Nada, Un poco, Medianamente, Mucho).",
        "evaluar": "Puntuación de 0 a 3 puntos para el índice de calidad de vida."
    },
    {
        "file": "32_quiz_q8_sociales.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 8 — Afectación en Relaciones Sociales",
        "route": "Modal Quiz · Paso 8 (Subpregunta)",
        "type": "Escala 0 a 3",
        "desc": "Se activa si seleccionó 'Vida social' en Q6. Mide el aislamiento, vergüenza o rechazo en reuniones familiares y amigos.",
        "evaluar": "Puntuación y tono empático sobre el aislamiento social."
    },
    {
        "file": "33_quiz_q8b_sexual.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 8b — Afectación en Vida Sexual",
        "route": "Modal Quiz · Paso 8b (Subpregunta)",
        "type": "Escala 0 a 3",
        "desc": "Se activa si seleccionó 'Vida sexual' en Q6. Aborda el impacto en la intimidad con la pareja con tacto y respeto.",
        "evaluar": "Tratamiento adecuado y discreto de una dimensión clínica sensible."
    },
    {
        "file": "34_quiz_q9_vestimenta.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 9 — Influencia en la Ropa que Usas",
        "route": "Modal Quiz · Paso 9 (Subpregunta)",
        "type": "Escala 0 a 3",
        "desc": "Se activa si seleccionó 'Vestimenta' en Q6. Evalúa si el paciente evita telas, elásticos o ropa ajustada por miedo al roce o calor.",
        "evaluar": "Puntuación en la escala global de calidad de vida."
    },
    {
        "file": "35_quiz_q10_trabajo.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 10 — Afectación en Trabajo / Estudios",
        "route": "Modal Quiz · Paso 10 (Subpregunta)",
        "type": "Escala 0 a 3",
        "desc": "Se activa si seleccionó 'Trabajo' en Q6. Mide falta de concentración, bajas laborales o baja productividad por prurito constante.",
        "evaluar": "Relevancia del dato socioeconómico y laboral para el reporte."
    },
    {
        "file": "36_quiz_q11_sueno.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 11 — Afectación en Calidad de Sueño",
        "route": "Modal Quiz · Paso 11 (Subpregunta)",
        "type": "Escala 0 a 3",
        "desc": "Se activa si seleccionó 'Sueño' en Q6. Mide despertares nocturnos causados por comezón nocturna intensa.",
        "evaluar": "Conexión con los contenidos de madrugada del sitio web."
    },
    {
        "file": "37_quiz_q12_detonantes.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 12 — Posibles Desencadenantes",
        "route": "Modal Quiz · Paso 12 (Multi-select)",
        "type": "Contexto & Factores",
        "desc": "Identificación de factores: Calor/duchas, Frío, Estrés emocional, Sudor/ejercicio, Ropa ajustada o 'Sin detonante claro' (espontáneo).",
        "evaluar": "Distinción clave: la UCE se caracteriza por no requerir detonante externo para brotar."
    },
    {
        "file": "38_quiz_q13_horario.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 13 — Momento de Mayor Intensidad",
        "route": "Modal Quiz · Paso 13",
        "type": "Patrón Circadiano",
        "desc": "Pregunta si los brotes ocurren en la mañana, tarde, noche/madrugada o sin patrón horario establecido.",
        "evaluar": "Utilidad de este dato para la redacción de la Sección 4 del reporte clínico."
    },
    {
        "file": "39_quiz_q14_consulta_previa.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 14 — Consulta Médica Previa",
        "route": "Modal Quiz · Paso 14",
        "type": "Historial de Atención",
        "desc": "¿Has consultado previamente a un médico sobre estos síntomas? (Sí / No). Si responde 'No', se saltan los pasos 15 y 16.",
        "evaluar": "Bifurcación lógica hacia el historial médico."
    },
    {
        "file": "40_quiz_q15_frecuencia_consulta.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 15 — Frecuencia de Consulta Médica",
        "route": "Modal Quiz · Paso 15 (Condicional)",
        "type": "Historial de Atención",
        "desc": "Se activa si consultó previamente: Pocas veces (1-2 veces), Frecuentemente (mensual) o En seguimiento continuo con especialista.",
        "evaluar": "Ponderación del nivel de atención médica recibido."
    },
    {
        "file": "41_quiz_q16_especialista.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 16 — Especialistas Consultados",
        "route": "Modal Quiz · Paso 16 (Multi-select)",
        "type": "Especialidades Médicas",
        "desc": "Tipo de médico visitado: Médico general, Dermatólogo/a, Alergólogo/a.",
        "evaluar": "Relevancia para orientar al paciente hacia el especialista certificado correcto."
    },
    {
        "file": "42_quiz_q17_tratamiento_previo.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 17 — Tratamiento Previo",
        "route": "Modal Quiz · Paso 17",
        "type": "Historial Terapéutico",
        "desc": "¿Has tomado o aplicado algún tratamiento? (Sí / No). Si responde 'No', se saltan los pasos 18 y 19.",
        "evaluar": "Bifurcación hacia el tipo y efectividad del tratamiento."
    },
    {
        "file": "43_quiz_q18_tipo_tratamiento.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 18 — Tipo de Tratamiento Recibido",
        "route": "Modal Quiz · Paso 18 (Multi-select)",
        "type": "Escalón Terapéutico",
        "desc": "Nivel de tratamiento: Básico (antihistamínicos venta libre/remedios), Intermedio (antihistamínicos recetados/corticoides) o Avanzado (dosis altas/biológicos).",
        "evaluar": "Mapeo de los escalones de las guías internacionales de urticaria."
    },
    {
        "file": "44_quiz_q19_efectividad.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Quiz Paso 19 — Efectividad Percibida",
        "route": "Modal Quiz · Paso 19",
        "type": "Respuesta al Tratamiento",
        "desc": "¿Qué tan efectivo ha sido?: Nada efectivo, Poco efectivo, Medianamente efectivo o Muy efectivo. Clave para la llamada a la acción final.",
        "evaluar": "Disparador de la alerta de escalación y del banner hacia el Directorio Médico."
    },
    {
        "file": "45_quiz_resultado_reporte_completo.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Pantalla de Resultados — Reporte Clínico Completo",
        "route": "Modal Quiz · Pantalla Final",
        "type": "Reporte Clínico Estructurado",
        "desc": "Muestra: 1) Diagnóstico probabilístico (Puntuación Ben 0-35), 2) Reporte estructurado en 6 secciones, 3) Pregunta clave sugerida para la consulta, 4) Botones de Descargar PDF y Copiar, 5) Banner destacado al Directorio Médico.",
        "evaluar": "Aclaración de que es orientativo, diseño del bloque de reporte, claridad del CTA al directorio médico."
    },
    {
        "file": "23_quiz_resultado_asintomatico.png",
        "section": "SECCIÓN 4: QUIZ ORIENTATIVO",
        "title": "Pantalla de Resultados — Asintomático / Poco Probable",
        "route": "Modal Quiz · Resultado Ruta Corta",
        "type": "Resultado Alternativo",
        "desc": "Generada cuando el usuario marca 'Ninguno de estos' en el Paso 1. Informa que no presenta síntomas sugestivos de UCE y brinda recursos educativos.",
        "evaluar": "Tono tranquilizador y adecuado para usuarios sin sospecha diagnóstica."
    },

    # =========================================================
    # SECCIÓN 5: DIRECTORIO MÉDICO Y BLOG EDITORIAL
    # =========================================================
    {
        "file": "13_directorio_landing.png",
        "section": "SECCIÓN 5: DIRECTORIO & BLOG",
        "title": "Directorio Médico de Especialistas (Landing Principal)",
        "route": "URL: /directorio-medico/",
        "type": "Landing Principal con Buscador",
        "desc": "Página completa dedicada a la localización de dermatólogos y alergólogos certificados en México. Buscador por código postal, estado y especialidad.",
        "evaluar": "Validación médica de los perfiles mostrados, fuentes de información y claridad de llamada directa a consulta."
    },
    {
        "file": "14_blog_feed.png",
        "section": "SECCIÓN 5: DIRECTORIO & BLOG",
        "title": "Blog Editorial — Catálogo de Temas",
        "route": "URL: /blog/",
        "type": "Feed Editorial",
        "desc": "Diseño inspirado en plataformas editoriales (estilo Medium) con artículos filtrables por categoría, estado de 'Borrador IA en revisión' y tiempos de lectura.",
        "evaluar": "Etiquetado transparente de contenido pendiente de revisión médica/editorial."
    },
    {
        "file": "15_blog_articulo_1.png",
        "section": "SECCIÓN 5: DIRECTORIO & BLOG",
        "title": "Artículo Blog 1: ¿Qué significa UCE?",
        "route": "URL: /blog/que-significa-uce/",
        "type": "Artículo de Lectura",
        "desc": "Lectura educativa de 3 minutos sobre el significado de las siglas UCE, causas no alérgicas y qué preguntas llevar a la consulta médica.",
        "evaluar": "Citas a guías médicas internacionales (EAACI/GA²LEN) y lenguaje accesible."
    },
    {
        "file": "16_blog_articulo_2.png",
        "section": "SECCIÓN 5: DIRECTORIO & BLOG",
        "title": "Artículo Blog 2: Guía para preparar tu consulta",
        "route": "URL: /blog/preparar-consulta/",
        "type": "Artículo de Lectura",
        "desc": "Consejos prácticos para que el paciente anote sus observaciones, cómo registrar el impacto en su vida y enlace al Diario de Síntomas.",
        "evaluar": "Recomendaciones prácticas y empáticas para pacientes de primera consulta."
    },
    {
        "file": "17_blog_articulo_3.png",
        "section": "SECCIÓN 5: DIRECTORIO & BLOG",
        "title": "Artículo Blog 3: Mitos de las dietas restrictivas en UCE",
        "route": "URL: /blog/dietas-restrictivas-uce/",
        "type": "Artículo de Lectura",
        "desc": "Desmitifica la suspensión innecesaria de alimentos sin diagnóstico alérgico confirmado y orienta a buscar diagnóstico profesional.",
        "evaluar": "Aclaración médica rigurosa sobre por qué la UCE no es una alergia alimentaria común."
    }
]

def create_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]  # Blank slide

    # -------------------------------------------------------------
    # SLIDE 1: PORTADA EJECUTIVA
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    
    bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_INK
    bg.line.fill.background()

    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.2))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_ACCENT
    accent_bar.line.fill.background()

    tx_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.3), Inches(11.0), Inches(3.2))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "NOVARTIS MÉXICO · MÁS QUE URTICARIA"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT

    p1 = tf.add_paragraph()
    p1.text = "Catálogo Maestro de Pantallas & Workflow"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_before = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = "Auditoría visual completa de 45 pantallas (Home, Landings, Quiz, Urgencia 911, Directorio y Blog)\ncon cajas de comentarios integradas para revisión del equipo de Novartis y Edward Alarcón."
    p2.font.size = Pt(17)
    p2.font.color.rgb = RGBColor(200, 215, 220)
    p2.space_before = Pt(12)

    card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.7), Inches(10.9), Inches(2.1))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(44, 72, 80)
    card.line.color.rgb = COLOR_ACCENT
    card.line.width = Pt(1.5)

    c_tf = card.text_frame
    c_tf.word_wrap = True
    cp1 = c_tf.paragraphs[0]
    cp1.text = "📝 GUÍA PARA REVISIÓN Y COMENTARIOS:"
    cp1.font.size = Pt(13)
    cp1.font.bold = True
    cp1.font.color.rgb = COLOR_WHITE

    cp2 = c_tf.add_paragraph()
    cp2.text = "• Cada diapositiva contiene una captura en alta definición de la pantalla completa del sitio o paso del Quiz.\n• En la columna derecha dispones de una caja de texto editable donde puedes escribir tus notas, correcciones de copy o sugerencias de diseño.\n• Puedes marcar el estado de cada vista: [Aprobado] [Con observaciones] o [Pendiente médico].\n• Archivo local .PPTX editable en Microsoft PowerPoint sin necesidad de conexión externa."
    cp2.font.size = Pt(11)
    cp2.font.color.rgb = RGBColor(220, 230, 235)
    cp2.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 2: ÍNDICE GENERAL / WORKFLOW
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    
    s2_head = slide2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(1.1))
    s2_tf = s2_head.text_frame
    p_s2_0 = s2_tf.paragraphs[0]
    p_s2_0.text = "ESTRUCTURA DEL SITIO · ÍNDICE DE SECCIONES"
    p_s2_0.font.size = Pt(12)
    p_s2_0.font.bold = True
    p_s2_0.font.color.rgb = COLOR_ACCENT

    p_s2_1 = s2_tf.add_paragraph()
    p_s2_1.text = "Índice General de Pantallas Documentadas (45 Pantallas)"
    p_s2_1.font.size = Pt(24)
    p_s2_1.font.bold = True
    p_s2_1.font.color.rgb = COLOR_INK

    sections_info = [
        ("SECCIÓN 1", "Home & Variantes de Contexto", "5 Pantallas", "Vistas de inicio dinámicas según clima y hora (Día, Calor, Frío, Madrugada y Noche cálida).", COLOR_ACCENT),
        ("SECCIÓN 2", "Landings Informativas y Clínicas", "4 Pantallas", "Síntomas detallados, Opciones de Tratamiento, Sobre Nosotros (Acerca de) y Error 404.", COLOR_TEAL),
        ("SECCIÓN 3", "Recursos & Herramientas", "4 Pantallas", "Hub de Recursos, Diario de Síntomas interactivo, Testimonios y Ficha de tono de piel.", COLOR_MUTED),
        ("SECCIÓN 4", "Quiz Orientativo & Urgencia 911", "27 Pantallas", "Flujo completo de evaluación clínica (Q0 a Q19), Pantalla de Urgencia 911, Directorio embebido y Resultados.", COLOR_ACCENT),
        ("SECCIÓN 5", "Directorio Médico & Blog", "5 Pantallas", "Buscador de especialistas certificados, Catálogo editorial y 3 artículos de lectura.", COLOR_TEAL)
    ]

    card_width = Inches(2.25)
    card_height = Inches(5.1)
    card_gap = Inches(0.18)
    start_x = Inches(0.8)
    start_y = Inches(1.65)

    for i, (sec_code, sec_title, sec_count, sec_desc, sec_color) in enumerate(sections_info):
        cx = start_x + i * (card_width + card_gap)
        scard = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, start_y, card_width, card_height)
        scard.fill.solid()
        scard.fill.fore_color.rgb = COLOR_BG_CARD
        scard.line.color.rgb = COLOR_LINE
        scard.line.width = Pt(1)

        stf = scard.text_frame
        stf.word_wrap = True
        
        sp0 = stf.paragraphs[0]
        sp0.text = sec_code
        sp0.font.size = Pt(11)
        sp0.font.bold = True
        sp0.font.color.rgb = sec_color

        sp1 = stf.add_paragraph()
        sp1.text = sec_title
        sp1.font.size = Pt(13)
        sp1.font.bold = True
        sp1.font.color.rgb = COLOR_INK
        sp1.space_before = Pt(6)

        sp2 = stf.add_paragraph()
        sp2.text = sec_count
        sp2.font.size = Pt(12)
        sp2.font.bold = True
        sp2.font.color.rgb = sec_color
        sp2.space_before = Pt(8)

        sp3 = stf.add_paragraph()
        sp3.text = sec_desc
        sp3.font.size = Pt(9.5)
        sp3.font.color.rgb = COLOR_MUTED
        sp3.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDES DE PANTALLAS (1 POR CADA CAPTURA)
    # -------------------------------------------------------------
    screenshot_dir = os.path.abspath("screenshots")
    total_screens = len(SCREENS_ORDERED)

    for idx, screen in enumerate(SCREENS_ORDERED, start=1):
        img_path = os.path.join(screenshot_dir, screen["file"])
        if not os.path.exists(img_path):
            print(f"Advertencia: no se encontró {img_path}, omitiendo slide...")
            continue

        slide = prs.slides.add_slide(blank_layout)

        # Header Bar
        hdr_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(12.13), Inches(0.85))
        htf = hdr_box.text_frame
        htf.word_wrap = True

        hp0 = htf.paragraphs[0]
        hp0.text = f"{screen['section']} · PANTALLA {idx:02d} DE {total_screens:02d}"
        hp0.font.size = Pt(10)
        hp0.font.bold = True
        hp0.font.color.rgb = COLOR_ACCENT

        hp1 = htf.add_paragraph()
        hp1.text = screen["title"]
        hp1.font.size = Pt(17)
        hp1.font.bold = True
        hp1.font.color.rgb = COLOR_INK

        hp2 = htf.add_paragraph()
        hp2.text = f"{screen['route']}  |  Formato: {screen['type']}"
        hp2.font.size = Pt(9.5)
        hp2.font.color.rgb = COLOR_MUTED

        # COLUMNA IZQUIERDA: MARCO Y CAPTURA ESCALADA PROPORCIONALMENTE
        img_frame_x = Inches(0.6)
        img_frame_y = Inches(1.2)
        max_frame_w = Inches(7.8)
        max_frame_h = Inches(5.9)

        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_frame_x, img_frame_y, max_frame_w, max_frame_h)
        frame.fill.solid()
        frame.fill.fore_color.rgb = COLOR_BG_CARD
        frame.line.color.rgb = COLOR_LINE
        frame.line.width = Pt(1)

        # Cálculo de aspect ratio con Pillow
        with Image.open(img_path) as im:
            orig_w, orig_h = im.size

        avail_w = max_frame_w - Inches(0.2)
        avail_h = max_frame_h - Inches(0.2)
        img_aspect = orig_w / orig_h
        frame_aspect = avail_w / avail_h

        if img_aspect > frame_aspect:
            final_w = avail_w
            final_h = avail_w / img_aspect
        else:
            final_h = avail_h
            final_w = avail_h * img_aspect

        # Centrado
        pos_x = img_frame_x + (max_frame_w - final_w) / 2
        pos_y = img_frame_y + (max_frame_h - final_h) / 2

        slide.shapes.add_picture(img_path, pos_x, pos_y, width=final_w, height=final_h)

        # COLUMNA DERECHA: TARJETA DE AUDITORÍA Y COMENTARIOS
        card_x = Inches(8.55)
        card_y = Inches(1.2)
        card_w = Inches(4.18)

        # 1. Tarjeta informativa
        info_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_w, Inches(2.75))
        info_card.fill.solid()
        info_card.fill.fore_color.rgb = COLOR_WHITE
        info_card.line.color.rgb = COLOR_LINE
        info_card.line.width = Pt(1)

        itf = info_card.text_frame
        itf.word_wrap = True

        ip0 = itf.paragraphs[0]
        ip0.text = "🎯 Propósito de la pantalla:"
        ip0.font.size = Pt(11)
        ip0.font.bold = True
        ip0.font.color.rgb = COLOR_TEAL

        ip1 = itf.add_paragraph()
        ip1.text = screen["desc"]
        ip1.font.size = Pt(9.5)
        ip1.font.color.rgb = COLOR_INK
        ip1.space_before = Pt(4)

        ip2 = itf.add_paragraph()
        ip2.text = "🔍 Aspectos clave a evaluar:"
        ip2.font.size = Pt(10.5)
        ip2.font.bold = True
        ip2.font.color.rgb = COLOR_ACCENT
        ip2.space_before = Pt(6)

        ip3 = itf.add_paragraph()
        ip3.text = screen["evaluar"]
        ip3.font.size = Pt(9)
        ip3.font.color.rgb = COLOR_MUTED
        ip3.space_before = Pt(3)

        # 2. Caja editable para comentarios del cliente
        comm_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y + Inches(2.9), card_w, Inches(3.0))
        comm_card.fill.solid()
        comm_card.fill.fore_color.rgb = RGBColor(255, 253, 248) # tono crema editable
        comm_card.line.color.rgb = COLOR_ACCENT
        comm_card.line.width = Pt(1.5)

        ctf = comm_card.text_frame
        ctf.word_wrap = True

        cp0 = ctf.paragraphs[0]
        cp0.text = "✍️ ESPACIO PARA COMENTARIOS (EDWARD / NOVARTIS):"
        cp0.font.size = Pt(10)
        cp0.font.bold = True
        cp0.font.color.rgb = COLOR_ACCENT

        cp1 = ctf.add_paragraph()
        cp1.text = "Estado: [   ] Aprobado   [   ] Con observaciones   [   ] Pendiente médico"
        cp1.font.size = Pt(9)
        cp1.font.bold = True
        cp1.font.color.rgb = COLOR_INK
        cp1.space_before = Pt(4)

        cp2 = ctf.add_paragraph()
        cp2.text = "Escribe aquí los comentarios, ajustes de copy o cambios solicitados para esta pantalla:\n\n"
        cp2.font.size = Pt(9)
        cp2.font.color.rgb = COLOR_MUTED
        cp2.space_before = Pt(6)

    output_filename = "Novartis_Mas_Que_Urticaria_Catalogo_Pantallas.pptx"
    prs.save(output_filename)
    print(f"Presentacion generada con exito: {os.path.abspath(output_filename)}")

if __name__ == "__main__":
    create_presentation()
